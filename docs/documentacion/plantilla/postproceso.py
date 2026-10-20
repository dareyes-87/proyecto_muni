"""Ajustes al .docx que genera pandoc.

Uso: python postproceso.py <entrada.docx> <salida.docx>

- Vuelve a aplicar los estilos de estilos.py (pandoc agrega algunos propios).
- Coloca el índice general (--toc) después de la portada, en @@INDICE_GENERAL@@.
- Divide el documento en dos secciones en @@INICIO_CUERPO@@: preliminares en
  romanos (i, ii…) y cuerpo en arábigos desde 1.
- "CAPÍTULO N" lleva el salto de página; el título que le sigue, no.
- Texto de las tablas con el estilo "Texto de tabla" y fila de encabezado en negrita.
- Pide a Word actualizar los campos (índices) al abrir.
"""
import copy
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import estilos

MARCA_INDICE = "@@INDICE_GENERAL@@"
MARCA_CUERPO = "@@INICIO_CUERPO@@"


def parrafo_con_texto(doc, texto):
    for p in doc.paragraphs:
        if p.text.strip() == texto:
            return p
    return None


def mover_indice(doc):
    cuerpo = doc.element.body
    sdt = cuerpo.find(qn("w:sdt"))
    marca = parrafo_con_texto(doc, MARCA_INDICE)
    if marca is None:
        return
    if sdt is not None:
        marca._p.addprevious(sdt)
        # Marcar el campo TOC como pendiente de actualizar.
        for fc in sdt.iter(qn("w:fldChar")):
            if fc.get(qn("w:fldCharType")) == "begin":
                fc.set(qn("w:dirty"), "true")
    cuerpo.remove(marca._p)


def dividir_secciones(doc):
    marca = parrafo_con_texto(doc, MARCA_CUERPO)
    if marca is None:
        return
    cuerpo = doc.element.body
    sect_final = cuerpo.find(qn("w:sectPr"))
    previo = marca._p.getprevious()
    if previo is None:          # no hay preliminares
        cuerpo.remove(marca._p)
        return
    # Sección de preliminares: copia de la final, numeración en romanos.
    sect_prelim = copy.deepcopy(sect_final)
    for e in sect_prelim.findall(qn("w:pgNumType")):
        sect_prelim.remove(e)
    num = OxmlElement("w:pgNumType")
    num.set(qn("w:fmt"), "lowerRoman")
    num.set(qn("w:start"), "1")
    sect_prelim.append(num)
    ppr = marca._p.get_or_add_pPr()
    for hijo in list(marca._p):
        if hijo is not ppr:
            marca._p.remove(hijo)
    ppr.append(sect_prelim)
    # Cuerpo: arábigos desde 1.
    for e in sect_final.findall(qn("w:pgNumType")):
        sect_final.remove(e)
    num = OxmlElement("w:pgNumType")
    num.set(qn("w:fmt"), "decimal")
    num.set(qn("w:start"), "1")
    sect_final.append(num)
    # Orden de elementos exigido por el esquema de sectPr: pgNumType antes de cols/titlePg/docGrid.
    for sect in (sect_prelim, sect_final):
        n = sect.find(qn("w:pgNumType"))
        for nombre in ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                       "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid"):
            otro = sect.find(qn(nombre))
            if otro is not None:
                otro.addprevious(n)
                break


def titulos_de_capitulo(doc):
    anterior = None
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        es_h1 = p.style is not None and p.style.name == "Heading 1"
        if es_h1 and anterior is not None and anterior.text.strip().upper().startswith("CAPÍTULO"):
            p.paragraph_format.page_break_before = False
        anterior = p if es_h1 else None


def tablas(doc):
    estilo = doc.styles["Texto de tabla"]
    for t in doc.tables:
        for i, fila in enumerate(t.rows):
            encabezado = fila._tr.trPr is not None and fila._tr.trPr.find(qn("w:tblHeader")) is not None
            for celda in fila.cells:
                for p in celda.paragraphs:
                    p.style = estilo
                    if i == 0 and encabezado:
                        for r in p.runs:
                            r.bold = True


def actualizar_campos_al_abrir(doc):
    settings = doc.settings.element
    for e in settings.findall(qn("w:updateFields")):
        settings.remove(e)
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def main():
    doc = Document(sys.argv[1])
    estilos.aplicar(doc)
    mover_indice(doc)
    titulos_de_capitulo(doc)
    tablas(doc)
    dividir_secciones(doc)
    actualizar_campos_al_abrir(doc)
    doc.save(sys.argv[2])


if __name__ == "__main__":
    main()
