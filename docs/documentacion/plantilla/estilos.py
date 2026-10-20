"""Estilos de la tesis (00_lineamientos.md, "Formato del documento").

Se usa de dos formas:
  python estilos.py <default_reference.docx> <referencia.docx>
      Crea plantilla/referencia.docx a partir del reference.docx por defecto de pandoc.
  import estilos; estilos.aplicar(doc)
      postproceso.py lo vuelve a aplicar sobre el .docx generado, por si pandoc
      agrega estilos propios (p. ej. "Source Code") que no toma de la referencia.
"""
import sys
import warnings

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FUENTE = "Times New Roman"
FUENTE_CODIGO = "Courier New"
TAM_TABLA = 11          # se permite de 10 a 12 pt dentro de tablas
SANGRIA = Cm(1.27)
MARGEN = Cm(2.54)

# "Caption" se busca por su id; python-docx avisa pero lo resuelve bien.
warnings.filterwarnings("ignore", "style lookup by style_id")


def _fuente(estilo, nombre=FUENTE, tam=None, negrita=None, cursiva=False, color=True):
    """Fija la fuente en los cuatro atributos de rFonts y quita las fuentes del tema."""
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for atr in list(rfonts.attrib):
        if atr.endswith("Theme"):
            del rfonts.attrib[atr]
    for atr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atr), nombre)
    if tam is not None:
        estilo.font.size = Pt(tam)
    if negrita is not None:
        estilo.font.bold = negrita
    estilo.font.italic = cursiva
    if color:
        estilo.font.color.rgb = RGBColor(0, 0, 0)


def _parrafo(estilo, alineacion, interlineado, sangria=Cm(0), izquierda=Cm(0)):
    pf = estilo.paragraph_format
    pf.alignment = alineacion
    pf.line_spacing_rule = interlineado
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = izquierda
    pf.first_line_indent = sangria


def _estilo(doc, nombre, base="Normal"):
    try:
        estilo = doc.styles[nombre]
    except KeyError:
        estilo = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
    estilo.base_style = doc.styles[base]
    estilo.quick_style = True
    return estilo


def _sin_nivel_esquema(estilo):
    ppr = estilo.element.get_or_add_pPr()
    for e in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(e)
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "9")
    ppr.append(lvl)


def _bordes_tabla(estilo):
    """Líneas horizontales (superior, inferior y entre filas), sin verticales."""
    tblpr = estilo.element.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        estilo.element.append(tblpr)
    for e in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(e)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "insideH"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        bordes.append(b)
    tblpr.append(bordes)


def aplicar(doc):
    E = doc.styles
    J, I, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
    DOBLE, SENCILLO = WD_LINE_SPACING.DOUBLE, WD_LINE_SPACING.SINGLE

    # Página: Carta, 2.54 cm en los cuatro lados.
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = MARGEN

    # Cuerpo del texto.
    for nombre in ("Normal", "Body Text", "First Paragraph"):
        _fuente(E[nombre], tam=12)
        _parrafo(E[nombre], J, DOBLE, SANGRIA)
    # Listas compactas de pandoc: como el cuerpo pero sin sangría de primera línea.
    _fuente(E["Compact"], tam=12)
    _parrafo(E["Compact"], J, DOBLE)
    _fuente(E["Block Text"], tam=12, color=False)
    _parrafo(E["Block Text"], J, DOBLE, izquierda=SANGRIA)

    # Títulos: TNR 12 negrita. H1 centrado, mayúsculas, salto de página antes.
    for n in range(1, 10):
        h = E[f"Heading {n}"]
        _fuente(h, tam=12, negrita=True)
        _parrafo(h, C if n == 1 else I, DOBLE)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = n == 1
        h.font.all_caps = n == 1
    for nombre in ("Title", "Subtitle", "Author", "Date", "Abstract Title", "TOC Heading"):
        _fuente(E[nombre], tam=12, negrita=nombre in ("Title", "TOC Heading", "Abstract Title"))
        _parrafo(E[nombre], C, DOBLE)
    E["TOC Heading"].font.all_caps = True
    E["TOC Heading"].paragraph_format.page_break_before = True
    _sin_nivel_esquema(E["TOC Heading"])

    tp = _estilo(doc, "Titulo preliminar")      # ÍNDICE DE FIGURAS / DE TABLAS
    _fuente(tp, tam=12, negrita=True)
    _parrafo(tp, C, DOBLE)
    tp.font.all_caps = True
    tp.paragraph_format.page_break_before = True
    tp.paragraph_format.keep_with_next = True
    _sin_nivel_esquema(tp)

    # Títulos y fuentes de figura y tabla: TNR 12, interlineado sencillo.
    for nombre in ("Caption", "Table Caption", "Image Caption"):
        _fuente(E[nombre], tam=12)
        _parrafo(E[nombre], I, SENCILLO)
    for nombre in ("Titulo de tabla", "Titulo de figura"):
        st = _estilo(doc, nombre, "Caption")
        _fuente(st, tam=12)
        _parrafo(st, I, SENCILLO)
        st.paragraph_format.keep_with_next = True
    fu = _estilo(doc, "Fuente", "Caption")
    _fuente(fu, tam=12)
    _parrafo(fu, I, SENCILLO)
    im = _estilo(doc, "Imagen")
    _parrafo(im, C, SENCILLO)
    im.paragraph_format.keep_with_next = True
    for nombre in ("Figure", "Captioned Figure"):
        _parrafo(E[nombre], C, SENCILLO)

    # Tablas.
    tt = _estilo(doc, "Texto de tabla")
    _fuente(tt, tam=TAM_TABLA)
    _parrafo(tt, I, SENCILLO)
    _fuente(E["Table"], tam=TAM_TABLA)
    _bordes_tabla(E["Table"])

    # Código: monoespaciada 10, sencillo, sin sangría.
    sc = _estilo(doc, "Source Code")
    _fuente(sc, FUENTE_CODIGO, tam=10, color=False)
    _parrafo(sc, I, SENCILLO)
    _fuente(E["Verbatim Char"], FUENTE_CODIGO, color=False)

    # Referencias: sangría francesa de 1.27 cm.
    bib = E["Bibliography"]
    _fuente(bib, tam=12)
    _parrafo(bib, I, DOBLE, sangria=-SANGRIA, izquierda=SANGRIA)

    # Notas al pie.
    _fuente(E["Footnote Text"], tam=10)
    _parrafo(E["Footnote Text"], J, SENCILLO)

    # Entradas de índices (Word los crea con base en Normal: doble y con sangría).
    for nombre, izq in (("toc 1", 0), ("toc 2", 0.5), ("toc 3", 1.0), ("table of figures", 0)):
        st = _estilo(doc, nombre)
        _fuente(st, tam=12)
        _parrafo(st, I, SENCILLO, izquierda=Cm(izq))
        st.paragraph_format.space_after = Pt(6)


def _pie_con_numero(doc):
    """Pie de página centrado con el campo PAGE."""
    pie = doc.sections[0].footer
    pie.is_linked_to_previous = False
    p = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    for tipo, texto in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        run = p.add_run()
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        elif texto == " PAGE ":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            run._r.append(it)
        else:
            run.text = texto


if __name__ == "__main__":
    origen, destino = sys.argv[1], sys.argv[2]
    doc = Document(origen)
    aplicar(doc)
    _pie_con_numero(doc)
    doc.save(destino)
    print(f"Plantilla creada: {destino}")
